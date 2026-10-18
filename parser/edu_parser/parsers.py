"""Extract text, tables and images from PDF and DOCX."""

import logging
import re
import statistics
import zipfile
from collections.abc import Iterator
from pathlib import Path

from .chunker import chunk_paragraphs, table_to_markdown
from .models import ParsedDocument, ParsedImage, ParsedTable

logger = logging.getLogger(__name__)

MIN_IMAGE_BYTES = 512
HEADING_FONT_RATIO = 1.25

_HEADING_REGEX = re.compile(r"^\d+(?:\.\d+)*\.?\s+[А-ЯA-Z]")

_HEADING_PREFIXES = (
    "пояснительная записка",
    "учебный план",
    "учебно-тематический план",
    "содержание программы",
    "оценочные материалы",
    "методические материалы",
    "методические рекомендации",
    "литература",
    "приложение",
    "кадровое обеспечение",
    "организационно-педагогические условия",
    "материально-технические условия",
    "учебно-методическое",
)


# ------------------------------------------------------------------------- DOCX
def _is_docx_heading(para) -> bool:
    style = getattr(para, "style", None)
    if style is None:
        return False
    sid = (getattr(style, "style_id", "") or "").lower()
    name = (getattr(style, "name", "") or "").lower()
    return any(needle in sid or needle in name for needle in ("heading", "заголовок", "title", "название"))


def _heading_level(para) -> int:
    sid = (getattr(para.style, "style_id", "") or "").lower()
    name = (getattr(para.style, "name", "") or "").lower()
    for s in (sid, name):
        for i in range(1, 10):
            if f"heading {i}" in s or f"заголовок {i}" in s or s.endswith(f" {i}"):
                return i
    return 0


def _looks_like_heading_fallback(text: str) -> bool:
    """Heuristic: numbered paragraph or a known section-title phrase."""
    if not text or len(text) > 200:
        return False
    if _HEADING_REGEX.match(text):
        return True
    lower = text.lower().rstrip(".")
    return any(lower.startswith(p) for p in _HEADING_PREFIXES)


def _fallback_heading_level(text: str) -> int:
    m = re.match(r"^(\d+(?:\.\d+)*)\.?", text)
    if m:
        return m.group(1).count(".") + 1
    return 1


def _iter_block_items(parent) -> Iterator:
    """Yield Paragraph / Table in document order."""
    from docx.document import Document as _Document
    from docx.oxml.table import CT_Tbl
    from docx.oxml.text.paragraph import CT_P
    from docx.table import Table, _Cell
    from docx.text.paragraph import Paragraph

    if isinstance(parent, _Document):
        parent_elm = parent.element.body
    elif isinstance(parent, _Cell):
        parent_elm = parent._tc
    else:
        raise ValueError(f"Unsupported parent: {type(parent)}")

    for child in parent_elm.iterchildren():
        if isinstance(child, CT_P):
            yield Paragraph(child, parent)
        elif isinstance(child, CT_Tbl):
            yield Table(child, parent)


def _docx_table_rows(table) -> list[list[str]]:
    rows: list[list[str]] = []
    for row in table.rows:
        cells = [(cell.text or "").strip() for cell in row.cells]
        if any(cells):
            rows.append(cells)
    return rows


def _ext_to_mime(ext: str) -> str:
    return {
        "png": "image/png",
        "jpg": "image/jpeg",
        "jpeg": "image/jpeg",
        "gif": "image/gif",
        "bmp": "image/bmp",
        "tif": "image/tiff",
        "tiff": "image/tiff",
        "emf": "image/emf",
        "wmf": "image/wmf",
        "svg": "image/svg+xml",
    }.get(ext, "application/octet-stream")


def _extract_docx_images(path: Path, doc_id: str, source_path: str) -> list[ParsedImage]:
    images: list[ParsedImage] = []
    try:
        with zipfile.ZipFile(path) as zf:
            ordinal = 0
            for name in zf.namelist():
                if not name.startswith("word/media/"):
                    continue
                data = zf.read(name)
                if len(data) < MIN_IMAGE_BYTES:
                    continue
                ext = name.rsplit(".", 1)[-1].lower()
                images.append(
                    ParsedImage(
                        doc_id=doc_id,
                        source_path=source_path,
                        page_number=None,
                        ordinal=ordinal,
                        data=data,
                        ext=ext,
                        mime=_ext_to_mime(ext),
                        metadata={"zip_name": name},
                    )
                )
                ordinal += 1
    except zipfile.BadZipFile:
        logger.warning("DOCX is not a valid zip: %s", path)
    return images


def parse_docx(path: Path, doc_id: str, source_path: str) -> ParsedDocument:
    """Parse a DOCX file into a ParsedDocument."""
    from docx import Document

    doc = Document(str(path))
    parsed = ParsedDocument(doc_id=doc_id, source_path=source_path, title=Path(path).stem)

    section_stack: dict[int, str] = {}
    para_buffer: list[str] = []
    ordinal_text = 0
    ordinal_table = 0

    def current_section() -> str | None:
        if not section_stack:
            return None
        levels = sorted(section_stack)
        return " > ".join(section_stack[lvl] for lvl in levels)

    def flush() -> None:
        nonlocal para_buffer, ordinal_text
        if not para_buffer:
            return
        chunks = chunk_paragraphs(
            paragraphs=para_buffer,
            doc_id=doc_id,
            source_path=source_path,
            section_title=current_section(),
            page_number=None,
            start_ordinal=ordinal_text,
        )
        parsed.text_chunks.extend(chunks)
        ordinal_text += len(chunks)
        para_buffer = []

    for block in _iter_block_items(doc):
        from docx.table import Table
        from docx.text.paragraph import Paragraph

        if isinstance(block, Paragraph):
            text = (block.text or "").strip()
            if not text:
                continue

            if _is_docx_heading(block):
                flush()
                level = _heading_level(block) or 1
                section_stack = {k: v for k, v in section_stack.items() if k < level}
                section_stack[level] = text
                continue

            if _looks_like_heading_fallback(text):
                flush()
                level = _fallback_heading_level(text)
                section_stack = {k: v for k, v in section_stack.items() if k < level}
                section_stack[level] = text
                continue

            para_buffer.append(text)

        elif isinstance(block, Table):
            rows = _docx_table_rows(block)
            if not rows:
                continue
            n_cols = max(len(r) for r in rows)
            para_buffer.append(f"[Таблица #{ordinal_table}: {len(rows)}×{n_cols}]")
            parsed.tables.append(
                ParsedTable(
                    doc_id=doc_id,
                    source_path=source_path,
                    section_title=current_section(),
                    page_number=None,
                    ordinal=ordinal_table,
                    rows=rows,
                    as_text=table_to_markdown(rows),
                )
            )
            ordinal_table += 1

    flush()
    parsed.images = _extract_docx_images(path, doc_id, source_path)
    return parsed


# -------------------------------------------------------------------------- PDF
class _TocTracker:
    """Tracks current section stack driven by a PDF TOC."""

    def __init__(self, toc: list[list]):
        self._toc = toc or []
        self._idx = 0
        self._stack: dict[int, str] = {}

    def advance_to_page(self, page_num: int) -> None:
        while self._idx < len(self._toc) and self._toc[self._idx][2] <= page_num:
            level, title, _ = self._toc[self._idx]
            level = max(1, int(level))
            self._stack = {k: v for k, v in self._stack.items() if k < level}
            self._stack[level] = title
            self._idx += 1

    @property
    def current(self) -> str | None:
        if not self._stack:
            return None
        return " > ".join(self._stack[lvl] for lvl in sorted(self._stack))


def _page_text_blocks(page) -> list[tuple[str, float]]:
    """Return [(block_text, max_font_size_in_block), ...]."""
    out: list[tuple[str, float]] = []
    data = page.get_text("dict")
    for block in data.get("blocks", []):
        if block.get("type") != 0:
            continue
        line_texts: list[str] = []
        max_size = 0.0
        for line in block.get("lines", []):
            parts: list[str] = []
            for span in line.get("spans", []):
                t = span.get("text", "")
                if t.strip():
                    parts.append(t)
                max_size = max(max_size, float(span.get("size", 0)))
            if parts:
                line_texts.append("".join(parts).rstrip())
        text = "\n".join(line_texts).strip()
        if text:
            out.append((text, max_size))
    return out


def _median_font_size(page) -> float:
    sizes: list[float] = []
    data = page.get_text("dict")
    for block in data.get("blocks", []):
        if block.get("type") != 0:
            continue
        for line in block.get("lines", []):
            for span in line.get("spans", []):
                if span.get("text", "").strip():
                    sizes.append(float(span.get("size", 0)))
    return statistics.median(sizes) if sizes else 10.0


def _extract_pdf_images(
    page, doc, doc_id: str, source_path: str, page_num: int, start_ordinal: int
) -> list[ParsedImage]:
    images: list[ParsedImage] = []
    ordinal = start_ordinal
    try:
        refs = page.get_images(full=True)
    except Exception:
        return images

    for ref in refs:
        xref = ref[0]
        try:
            info = doc.extract_image(xref)
        except Exception as exc:
            logger.debug("skip image xref=%s: %s", xref, exc)
            continue
        data = info.get("image") or b""
        if len(data) < MIN_IMAGE_BYTES:
            continue
        ext = (info.get("ext") or "png").lower()
        images.append(
            ParsedImage(
                doc_id=doc_id,
                source_path=source_path,
                page_number=page_num,
                ordinal=ordinal,
                data=data,
                ext=ext,
                mime=_ext_to_mime(ext),
                metadata={
                    "xref": xref,
                    "width": info.get("width"),
                    "height": info.get("height"),
                },
            )
        )
        ordinal += 1
    return images


def parse_pdf(path: Path, doc_id: str, source_path: str) -> ParsedDocument:
    """Parse a PDF file into a ParsedDocument."""
    import pymupdf

    doc = pymupdf.open(str(path))
    parsed = ParsedDocument(doc_id=doc_id, source_path=source_path, title=Path(path).stem)

    try:
        toc = doc.get_toc() or []
    except Exception:
        toc = []
    tracker = _TocTracker(toc) if toc else None

    ordinal_text = 0
    ordinal_table = 0
    ordinal_image = 0

    for page_num, page in enumerate(doc, start=1):
        if tracker is not None:
            tracker.advance_to_page(page_num)
        section_title = tracker.current if tracker else None

        page_paragraphs: list[str] = []
        median = _median_font_size(page) or 10.0
        heading_threshold = median * HEADING_FONT_RATIO

        for block_text, max_size in _page_text_blocks(page):
            is_heading_candidate = tracker is None and max_size >= heading_threshold and len(block_text) < 200
            if is_heading_candidate:
                if page_paragraphs:
                    chunks = chunk_paragraphs(
                        paragraphs=page_paragraphs,
                        doc_id=doc_id,
                        source_path=source_path,
                        section_title=section_title,
                        page_number=page_num,
                        start_ordinal=ordinal_text,
                    )
                    parsed.text_chunks.extend(chunks)
                    ordinal_text += len(chunks)
                    page_paragraphs = []
                section_title = block_text.strip() or section_title
            else:
                page_paragraphs.append(block_text)

        if page_paragraphs:
            chunks = chunk_paragraphs(
                paragraphs=page_paragraphs,
                doc_id=doc_id,
                source_path=source_path,
                section_title=section_title,
                page_number=page_num,
                start_ordinal=ordinal_text,
            )
            parsed.text_chunks.extend(chunks)
            ordinal_text += len(chunks)

        try:
            finder = page.find_tables()
            for tab in getattr(finder, "tables", []):
                rows = tab.extract()
                if not rows:
                    continue
                rows = [[("" if c is None else str(c)) for c in r] for r in rows]
                parsed.tables.append(
                    ParsedTable(
                        doc_id=doc_id,
                        source_path=source_path,
                        section_title=section_title,
                        page_number=page_num,
                        ordinal=ordinal_table,
                        rows=rows,
                        as_text=table_to_markdown(rows),
                    )
                )
                ordinal_table += 1
        except Exception as exc:
            logger.debug("table extraction failed on page %s: %s", page_num, exc)

        imgs = _extract_pdf_images(page, doc, doc_id, source_path, page_num, ordinal_image)
        parsed.images.extend(imgs)
        ordinal_image += len(imgs)

    doc.close()
    return parsed
