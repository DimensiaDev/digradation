"""Tests for the CLI entrypoint."""

import json

import pytest
from docx import Document

from edu_parser.cli import main as cli_main


def test_cli_help(capsys):
    with pytest.raises(SystemExit) as exc:
        cli_main(["--help"])
    assert exc.value.code == 0
    out = capsys.readouterr().out
    assert "Parse PDF and DOCX" in out
    assert "--input" in out
    assert "--output" in out


def test_cli_missing_input_returns_2(tmp_path, capsys):
    code = cli_main(["--input", str(tmp_path / "nope"), "--output", str(tmp_path / "out")])
    assert code == 2
    err = capsys.readouterr().err
    assert "not a directory" in err


def test_cli_parse_one_ok(tmp_path, capsys):
    src = tmp_path / "src"
    src.mkdir()
    d = Document()
    d.add_paragraph("Hello")
    d.save(src / "a.docx")

    code = cli_main(["--input", str(src), "--output", str(tmp_path / "out")])
    assert code == 0
    out = capsys.readouterr().out
    assert "Parsed: 1 OK" in out
    assert (tmp_path / "out" / "manifest.json").exists()


def test_cli_returns_1_on_failures(tmp_path):
    src = tmp_path / "src"
    src.mkdir()
    (src / "broken.pdf").write_bytes(b"junk")

    code = cli_main(["--input", str(src), "--output", str(tmp_path / "out")])
    assert code == 1


def test_cli_limit_respected(tmp_path):
    src = tmp_path / "src"
    src.mkdir()
    for name in ("a.docx", "b.docx"):
        d = Document()
        d.add_paragraph(f"Hello {name}")
        d.save(src / name)

    code = cli_main(
        [
            "--input",
            str(src),
            "--output",
            str(tmp_path / "out"),
            "--limit",
            "1",
        ]
    )
    assert code == 0
    m = json.loads((tmp_path / "out" / "manifest.json").read_text())
    assert m["total"] == 1


def test_cli_verbose_flag(tmp_path, caplog):
    src = tmp_path / "src"
    src.mkdir()
    d = Document()
    d.add_paragraph("Hello")
    d.save(src / "a.docx")

    code = cli_main(["--input", str(src), "--output", str(tmp_path / "out"), "-v"])
    assert code == 0


def test_cli_prune_flag(tmp_path):
    src = tmp_path / "src"
    src.mkdir()
    d = Document()
    d.add_paragraph("Hello")
    d.save(src / "a.docx")

    out = tmp_path / "out"
    out.mkdir()
    (out / "stale").mkdir()

    code = cli_main(["--input", str(src), "--output", str(out), "--prune"])
    assert code == 0
    assert not (out / "stale").exists()
