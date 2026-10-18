"""Smoke tests: chunker + end-to-end on a tiny fake DOCX."""

from pathlib import Path

from docx import Document

from edu_parser.chunker import chunk_paragraphs, table_to_markdown
from edu_parser.parsers import parse_docx
from edu_parser.pipeline import process_directory


def test_chunk_paragraphs_basic():
    paras = [f"Параграф номер {i} с каким-то текстом." for i in range(50)]
    chunks = chunk_paragraphs(
        paragraphs=paras,
        doc_id="d1",
        source_path="x.docx",
        section_title="Раздел 1",
        page_number=1,
        max_size=300,
        min_size=50,
        overlap=50,
    )
    assert len(chunks) > 1
    for c in chunks:
        assert c.doc_id == "d1"
        assert c.section_title == "Раздел 1"
        assert len(c.text) <= 400


def test_table_to_markdown():
    md = table_to_markdown([["Компетенция", "Индикатор"], ["УК-1", "..."]])
    assert "| Компетенция | Индикатор |" in md
    assert "| УК-1 | ... |" in md


def test_parse_docx_smoke(tmp_path: Path):
    src = tmp_path / "sample.docx"
    d = Document()
    d.add_heading("Раздел 1", level=1)
    d.add_paragraph("Текст первого раздела.")
    t = d.add_table(rows=2, cols=2)
    t.cell(0, 0).text = "Компетенция"
    t.cell(0, 1).text = "Индикатор"
    t.cell(1, 0).text = "УК-1"
    t.cell(1, 1).text = "Пример"
    d.save(src)

    doc = parse_docx(src, doc_id="d1", source_path="sample.docx")
    assert any(c.section_title and "Раздел 1" in c.section_title for c in doc.text_chunks)
    assert len(doc.tables) == 1
    assert doc.tables[0].rows[0][0] == "Компетенция"


def test_process_directory_isolates_errors(tmp_path: Path):
    src = tmp_path / "src"
    src.mkdir()
    d = Document()
    d.add_paragraph("Хелло")
    d.save(src / "ok.docx")
    (src / "broken.pdf").write_bytes(b"not a pdf at all")

    out = tmp_path / "out"
    results = process_directory(src, out)
    assert any(r["ok"] for r in results)
    assert any(not r["ok"] for r in results)
    assert (out / "manifest.json").exists()
