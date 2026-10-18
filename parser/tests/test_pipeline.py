"""Tests for the pipeline module: serialization, atomic save, pruning."""

from pathlib import Path

import pytest
from docx import Document

from edu_parser import pipeline
from edu_parser.models import ParsedDocument, ParsedImage, TextChunk
from edu_parser.pipeline import (
    _md_escape_inline,
    _prune_stale,
    _save,
    _serialize,
    _write_preview_md,
    process_directory,
)


# ---------------------------------------------------------------- helpers
def _make_docx(path: Path, text: str = "Hello") -> None:
    d = Document()
    d.add_paragraph(text)
    d.save(path)


# --------------------------------------------------------------- escaping
def test_md_escape_inline_plain():
    assert _md_escape_inline("no special chars") == "no special chars"


def test_md_escape_inline_specials():
    assert "\\*" in _md_escape_inline("a * b")
    assert "\\`" in _md_escape_inline("tick ` here")
    assert "\\[" in _md_escape_inline("[bracket]")


def test_write_preview_md_escapes_html(tmp_path: Path):
    doc = ParsedDocument(
        doc_id="d1",
        source_path="x.docx",
        title="Test",
        text_chunks=[
            TextChunk(
                doc_id="d1",
                source_path="x.docx",
                section_title="Section",
                page_number=1,
                ordinal=0,
                text="<script>alert(1)</script>",
            )
        ],
    )
    out = tmp_path / "preview.md"
    _write_preview_md(doc, out)
    content = out.read_text()
    assert "```text" in content
    assert "```" in content
    assert "## chunk #0" in content


# ------------------------------------------------------------- serialize
def test_serialize_without_images(tmp_path: Path):
    doc = ParsedDocument(
        doc_id="d1",
        source_path="x.docx",
        title="Test",
        text_chunks=[
            TextChunk(
                doc_id="d1",
                source_path="x.docx",
                section_title=None,
                page_number=None,
                ordinal=0,
                text="Hello",
            )
        ],
    )
    payload = _serialize(doc, images_dir=None)
    assert payload["doc_id"] == "d1"
    assert payload["counts"]["text_chunks"] == 1
    assert payload["images"] == []


def test_serialize_with_images(tmp_path: Path):
    doc = ParsedDocument(
        doc_id="d1",
        source_path="x.docx",
        title="Test",
        images=[
            ParsedImage(
                doc_id="d1",
                source_path="x.docx",
                page_number=1,
                ordinal=0,
                data=b"\x89PNG\r\n\x1a\n" + b"x" * 600,
                ext="png",
                mime="image/png",
            )
        ],
    )
    images_dir = tmp_path / "imgs"
    payload = _serialize(doc, images_dir=images_dir)
    assert payload["images"][0]["path"] is not None
    assert (tmp_path / "imgs" / "0000.png").exists()


# ------------------------------------------------------------------ save
def test_save_atomic_on_serialize_error(tmp_path, monkeypatch):
    doc = ParsedDocument(doc_id="d1", source_path="x.docx", title="Test")
    out = tmp_path / "out"
    out.mkdir()

    def boom(*args, **kwargs):
        raise RuntimeError("serialize exploded")

    monkeypatch.setattr(pipeline, "_serialize", boom)

    with pytest.raises(RuntimeError, match="serialize exploded"):
        _save(doc, out)

    assert not (out / "d1").exists()
    leftover = [p for p in out.iterdir() if p.name.startswith(".d1")]
    assert leftover == []


def test_save_writes_all_artifacts(tmp_path: Path):
    doc = ParsedDocument(
        doc_id="d1",
        source_path="x.docx",
        title="Test",
        text_chunks=[
            TextChunk(
                doc_id="d1",
                source_path="x.docx",
                section_title=None,
                page_number=None,
                ordinal=0,
                text="Hello",
            )
        ],
    )
    out = tmp_path / "out"
    out.mkdir()
    info = _save(doc, out)

    doc_dir = out / "d1"
    assert (doc_dir / "document.json").exists()
    assert (doc_dir / "preview.md").exists()
    assert info["counts"]["text_chunks"] == 1


# ----------------------------------------------------------------- prune
def test_prune_stale_removes_old(tmp_path: Path):
    out = tmp_path / "out"
    out.mkdir()
    (out / "keep").mkdir()
    (out / "stale").mkdir()
    (out / "manifest.json").write_text("{}")

    _prune_stale(out, {"keep"})

    assert (out / "keep").exists()
    assert not (out / "stale").exists()
    assert (out / "manifest.json").exists()


def test_process_directory_rejects_file_as_input(tmp_path: Path):
    not_dir = tmp_path / "file.txt"
    not_dir.write_text("x")
    with pytest.raises(NotADirectoryError):
        process_directory(not_dir, tmp_path / "out")


def test_process_directory_prune_stale_flag(tmp_path: Path):
    src = tmp_path / "src"
    src.mkdir()
    _make_docx(src / "a.docx")

    out = tmp_path / "out"
    out.mkdir()
    (out / "old").mkdir()

    results = process_directory(src, out, prune_stale=True)
    assert len(results) == 1
    assert results[0]["ok"] is True
    assert not (out / "old").exists()


# ---------------------------------------------------------- _iter_files coverage
def test_iter_files_skips_symlink(tmp_path: Path):
    """Symlink inside input root must be skipped (branch 35)."""
    from edu_parser.pipeline import _iter_files

    src = tmp_path / "src"
    src.mkdir()
    real = tmp_path / "outside.docx"
    _make_docx(real)
    link = src / "link.docx"
    link.symlink_to(real)

    files = list(_iter_files(src))
    assert all("link.docx" not in str(f) for f in files)


def test_iter_files_skips_unsupported(tmp_path: Path):
    """Files with non-supported extensions are ignored (branch 40/43)."""
    from edu_parser.pipeline import _iter_files

    src = tmp_path / "src"
    src.mkdir()
    (src / "notes.txt").write_text("x")
    (src / "image.png").write_bytes(b"x")
    _make_docx(src / "ok.docx")

    files = list(_iter_files(src))
    assert len(files) == 1
    assert files[0].name == "ok.docx"


# ----------------------------------------- _write_preview_md tables + metadata
def test_write_preview_md_with_tables(tmp_path: Path):
    """Cover the tables loop in _write_preview_md (lines 139-149)."""
    from edu_parser.models import ParsedTable

    doc = ParsedDocument(
        doc_id="d1",
        source_path="x.docx",
        title="Test",
        tables=[
            ParsedTable(
                doc_id="d1",
                source_path="x.docx",
                section_title="T1",
                page_number=2,
                ordinal=0,
                rows=[["A", "B"], ["C", "D"]],
                as_text="| A | B |\n| C | D |",
            )
        ],
    )
    out = tmp_path / "preview.md"
    _write_preview_md(doc, out)
    content = out.read_text()
    assert "## table #0" in content
    assert "page 2" in content
    assert "T1" in content
    assert "```markdown" in content
    assert "| A | B |" in content


def test_write_preview_md_chunk_without_section(tmp_path: Path):
    """Cover both branches of `if c.section_title` (92->97)."""
    doc = ParsedDocument(
        doc_id="d1",
        source_path="x.docx",
        title="Test",
        text_chunks=[
            TextChunk(
                doc_id="d1",
                source_path="x.docx",
                section_title=None,
                page_number=None,
                ordinal=0,
                text="no section",
            ),
            TextChunk(
                doc_id="d1",
                source_path="x.docx",
                section_title="Section",
                page_number=3,
                ordinal=1,
                text="with section",
            ),
        ],
    )
    out = tmp_path / "preview.md"
    _write_preview_md(doc, out)
    content = out.read_text()
    assert "## chunk #0" in content
    assert "page 3" in content
    assert "Section" in content
