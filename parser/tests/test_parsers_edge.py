"""Edge-case tests for parser helpers."""

import zipfile
from pathlib import Path

from docx import Document

from edu_parser import parsers
from edu_parser.parsers import (
    _ext_to_mime,
    _extract_docx_images,
    _fallback_heading_level,
    _looks_like_heading_fallback,
)


def test_looks_like_heading_fallback_numbered():
    assert _looks_like_heading_fallback("1. Introduction")
    assert _looks_like_heading_fallback("2.3. Methodology")


def test_looks_like_heading_fallback_known_prefix():
    assert _looks_like_heading_fallback("Литература")
    assert _looks_like_heading_fallback("Пояснительная записка")


def test_looks_like_heading_fallback_too_long():
    assert not _looks_like_heading_fallback("x" * 200)


def test_looks_like_heading_fallback_trailing_dot():
    assert not _looks_like_heading_fallback("1. This is a sentence.")


def test_looks_like_heading_fallback_empty():
    assert not _looks_like_heading_fallback("")


def test_fallback_heading_level_flat():
    assert _fallback_heading_level("1. Intro") == 1


def test_fallback_heading_level_nested():
    assert _fallback_heading_level("1.2. Intro") == 2
    assert _fallback_heading_level("1.2.3. Intro") == 3


def test_fallback_heading_level_unnumbered():
    assert _fallback_heading_level("Untitled") == 1


def test_ext_to_mime_known():
    assert _ext_to_mime("png") == "image/png"
    assert _ext_to_mime("jpeg") == "image/jpeg"
    assert _ext_to_mime("emf") == "image/emf"


def test_ext_to_mime_unknown():
    assert _ext_to_mime("xyz") == "application/octet-stream"


def test_extract_docx_images_bad_zip(tmp_path: Path):
    src = tmp_path / "bad.docx"
    src.write_bytes(b"not a zip")
    errors: list[str] = []
    imgs = _extract_docx_images(src, "d1", "bad.docx", errors)
    assert imgs == []
    assert len(errors) == 1
    assert "not a valid zip" in errors[0]


def test_extract_docx_images_oversized(tmp_path: Path, monkeypatch):
    monkeypatch.setattr(parsers, "MAX_IMAGE_BYTES_CAP", 1000)

    d = Document()
    d.add_paragraph("Text")
    src = tmp_path / "big.docx"
    d.save(src)

    with zipfile.ZipFile(src, "a") as zf:
        zf.writestr("word/media/big.png", b"x" * 2000)

    errors: list[str] = []
    imgs = _extract_docx_images(src, "d1", "big.docx", errors)
    assert imgs == []
    assert any("oversized" in e for e in errors)


def test_extract_docx_images_skips_tiny(tmp_path: Path):
    d = Document()
    d.add_paragraph("Text")
    src = tmp_path / "tiny.docx"
    d.save(src)

    with zipfile.ZipFile(src, "a") as zf:
        zf.writestr("word/media/tiny.png", b"x" * 10)

    errors: list[str] = []
    imgs = _extract_docx_images(src, "d1", "tiny.docx", errors)
    assert imgs == []


def test_extract_docx_images_valid(tmp_path: Path):
    d = Document()
    d.add_paragraph("Text")
    src = tmp_path / "ok.docx"
    d.save(src)

    payload = b"\x89PNG\r\n\x1a\n" + b"x" * 1000
    with zipfile.ZipFile(src, "a") as zf:
        zf.writestr("word/media/logo.png", payload)

    errors: list[str] = []
    imgs = _extract_docx_images(src, "d1", "ok.docx", errors)
    assert len(imgs) == 1
    assert imgs[0].ext == "png"
    assert imgs[0].mime == "image/png"
    assert errors == []
