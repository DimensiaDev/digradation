"""Targeted tests to close remaining coverage gaps."""

import zipfile
from pathlib import Path
from unittest.mock import MagicMock, patch

import pymupdf
import pytest
from docx import Document

from edu_parser.parsers import (
    _ext_to_mime,
    _extract_docx_images,
    _extract_pdf_images,
    _heading_level,
    _is_docx_heading,
    parse_docx,
    parse_pdf,
)
from edu_parser.pipeline import _prune_stale


# ---------------------------------------------------------------- _ext_to_mime
def test_ext_to_mime_covers_all_extensions():
    """Every entry in the _ext_to_mime mapping must be exercised."""
    for ext, mime in (
        ("png", "image/png"),
        ("jpg", "image/jpeg"),
        ("jpeg", "image/jpeg"),
        ("gif", "image/gif"),
        ("bmp", "image/bmp"),
        ("tif", "image/tiff"),
        ("tiff", "image/tiff"),
        ("emf", "image/emf"),
        ("wmf", "image/wmf"),
        ("svg", "image/svg+xml"),
    ):
        assert _ext_to_mime(ext) == mime


# ---------------------------------------------------------- _is_docx_heading
def test_is_docx_heading_plain_paragraph(tmp_path: Path):
    """Paragraph without style → False (line 43)."""
    d = Document()
    d.add_paragraph("plain")
    src = tmp_path / "x.docx"
    d.save(src)

    doc = Document(src)
    assert _is_docx_heading(doc.paragraphs[0]) is False


def test_heading_level_plain_paragraph(tmp_path: Path):
    """Non-heading paragraph → 0 (line 56)."""
    d = Document()
    d.add_paragraph("plain")
    src = tmp_path / "x.docx"
    d.save(src)

    doc = Document(src)
    assert _heading_level(doc.paragraphs[0]) == 0


# ------------------------------------------------- _extract_docx_images edge
def test_extract_docx_images_ignores_non_media(tmp_path: Path):
    """Non-media entries must be skipped (branch 237->210)."""
    d = Document()
    d.add_paragraph("hi")
    src = tmp_path / "x.docx"
    d.save(src)

    with zipfile.ZipFile(src, "a") as zf:
        zf.writestr("word/extra-notes.xml", b"<x/>")
        zf.writestr("word/media/logo.png", b"\x89PNG\r\n\x1a\n" + b"x" * 1000)

    errors: list[str] = []
    imgs = _extract_docx_images(src, "d1", "x.docx", errors)
    assert len(imgs) == 1
    assert errors == []


# -------------------------------------------------------- _extract_pdf_images
def test_extract_pdf_images_filters_and_accepts():
    """Cover MIN/MAX filter branches and successful append."""
    tiny = b"x" * 100
    ok = b"y" * 1000
    huge = b"z" * 3000

    page = MagicMock()
    page.get_images.return_value = [
        (1, 0, 0, 0, 0, 0, 0),
        (2, 0, 0, 0, 0, 0, 0),
        (3, 0, 0, 0, 0, 0, 0),
    ]

    doc = MagicMock()
    doc.extract_image.side_effect = [
        {"image": tiny, "ext": "png", "width": 1, "height": 1},
        {"image": ok, "ext": "png", "width": 10, "height": 10},
        {"image": huge, "ext": "png", "width": 100, "height": 100},
    ]

    with patch("edu_parser.parsers.MAX_IMAGE_BYTES_CAP", 2000):
        errors: list[str] = []
        images = _extract_pdf_images(page, doc, "d1", "x.pdf", 1, 0, errors)

    assert len(images) == 1
    assert images[0].data == ok
    assert images[0].page_number == 1
    assert any("oversized" in e for e in errors)


def test_extract_pdf_images_listing_error_is_recorded():
    """Cover the except around page.get_images."""
    page = MagicMock()
    page.get_images.side_effect = RuntimeError("corrupt image stream")

    errors: list[str] = []
    images = _extract_pdf_images(page, MagicMock(), "d1", "x.pdf", 1, 0, errors)

    assert images == []
    assert any("listing failed" in e for e in errors)


# ---------------------------------------------------------- parse_pdf branches
def _make_multi_page_pdf(path: Path, pages: list[str], toc: list[list] | None = None) -> None:
    doc = pymupdf.open()
    for text in pages:
        page = doc.new_page()
        page.insert_text((72, 72), text)
    if toc:
        doc.set_toc(toc)
    doc.save(str(path))
    doc.close()


def test_parse_pdf_with_heading_font(tmp_path: Path):
    """Trigger the is_heading_candidate branch in parse_pdf."""
    doc = pymupdf.open()
    page = doc.new_page()
    page.insert_text((72, 72), "BIG HEADING", fontsize=24)
    page.insert_text((72, 150), "Regular body text.", fontsize=10)
    src = tmp_path / "x.pdf"
    doc.save(str(src))
    doc.close()

    parsed = parse_pdf(src, doc_id="d1", source_path="x.pdf")
    assert any("BIG HEADING" in c.text for c in parsed.text_chunks)


def test_parse_pdf_empty_file_raises(tmp_path: Path):
    src = tmp_path / "empty.pdf"
    src.write_bytes(b"")
    with pytest.raises((pymupdf.FileDataError, pymupdf.EmptyFileError, ValueError)):
        parse_pdf(src, doc_id="d1", source_path="empty.pdf")


# --------------------------------------------------------- _prune_stale extras
def test_prune_stale_skips_non_dir(tmp_path: Path):
    """Regular file in output root must not be touched (branch)."""
    out = tmp_path / "out"
    out.mkdir()
    (out / "some.log").write_text("log")
    (out / "keep").mkdir()

    _prune_stale(out, {"keep"})
    assert (out / "some.log").exists()
    assert (out / "keep").exists()


# ------------------------------------------------ parse_docx empty paragraph
def test_parse_docx_empty_paragraph_skipped(tmp_path: Path):
    """Empty paragraph → continue (line 318)."""
    d = Document()
    d.add_paragraph("")
    d.add_paragraph("Content")
    src = tmp_path / "x.docx"
    d.save(src)

    parsed = parse_docx(src, doc_id="d1", source_path="x.docx")
    joined = "\n".join(c.text for c in parsed.text_chunks)
    assert "Content" in joined


# ---------------------------------------------- parse_docx table branch (349-353)
def test_parse_docx_with_table(tmp_path: Path):
    """Cover the Table branch in parse_docx."""
    d = Document()
    d.add_paragraph("Intro")
    t = d.add_table(rows=2, cols=2)
    t.cell(0, 0).text = "A"
    t.cell(0, 1).text = "B"
    t.cell(1, 0).text = "C"
    t.cell(1, 1).text = "D"
    d.add_paragraph("Outro")
    src = tmp_path / "x.docx"
    d.save(src)

    parsed = parse_docx(src, doc_id="d1", source_path="x.docx")
    assert len(parsed.tables) == 1
    assert parsed.tables[0].rows[0][0] == "A"
    assert parsed.tables[0].rows[1][1] == "D"


def test_parse_docx_table_marker_in_text(tmp_path: Path):
    """Table placeholder must appear in the surrounding text chunks."""
    d = Document()
    d.add_paragraph("Before")
    t = d.add_table(rows=1, cols=1)
    t.cell(0, 0).text = "X"
    d.add_paragraph("After")
    src = tmp_path / "x.docx"
    d.save(src)

    parsed = parse_docx(src, doc_id="d1", source_path="x.docx")
    joined = "\n".join(c.text for c in parsed.text_chunks)
    assert "[Table #0" in joined


def test_parse_docx_empty_table_skipped(tmp_path: Path):
    """Empty table → continue (branch 321->320)."""
    d = Document()
    d.add_paragraph("Intro")
    d.add_table(rows=1, cols=2)
    d.add_paragraph("Outro")
    src = tmp_path / "x.docx"
    d.save(src)

    parsed = parse_docx(src, doc_id="d1", source_path="x.docx")
    assert parsed.tables == []
    joined = "\n".join(c.text for c in parsed.text_chunks)
    assert "Intro" in joined
    assert "Outro" in joined


def test_extract_docx_images_unknown_extension(tmp_path: Path):
    """Unknown extension → application/octet-stream (fallback in _ext_to_mime)."""
    d = Document()
    d.add_paragraph("hi")
    src = tmp_path / "x.docx"
    d.save(src)

    with zipfile.ZipFile(src, "a") as zf:
        zf.writestr("word/media/data.xyz", b"\x89PNG\r\n\x1a\n" + b"x" * 1000)

    errors: list[str] = []
    imgs = _extract_docx_images(src, "d1", "x.docx", errors)
    assert len(imgs) == 1
    assert imgs[0].ext == "xyz"
    assert imgs[0].mime == "application/octet-stream"
