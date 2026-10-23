"""Tests for the DOCX renderer."""

import io
import time

from docx import Document

from app.agentic.models import (
    Competence,
    Curriculum,
    Hours,
    Module,
    Outcome,
    ProgramDraft,
    ProgramFormat,
    ProgramKind,
    ProgramMode,
    ProgramSummary,
    Schedule,
    Source,
    SourceKind,
    Topic,
)
from app.render import render_program_draft


# --------------------------------------------------------------------------- fixtures
def _hours(lecture: float = 0, practice: float = 0, lab: float = 0, self_study: float = 0) -> Hours:
    return Hours(lecture=lecture, practice=practice, lab=lab, self_study=self_study)


def _minimal_draft(**overrides) -> ProgramDraft:
    defaults = {
        "kind": ProgramKind.GENERAL_DEVELOPMENT,
        "title": "Тестовая программа",
        "summary": ProgramSummary(
            goal="Цель программы",
            target_audience="Взрослые слушатели",
            form="Очная",
            issued_document="Сертификат",
        ),
        "competencies": [
            Competence(
                id="comp-1",
                statement="Способен решать профессиональные задачи",
                knowledge=[Outcome(id="ro-k-1", statement="Базовые знания")],
            )
        ],
        "curriculum": Curriculum(
            modules=[
                Module(
                    id="m-1",
                    title="Модуль 1",
                    topics=[
                        Topic(
                            id="t-1.1",
                            title="Тема 1.1",
                            hours=_hours(lecture=2, practice=2),
                        )
                    ],
                )
            ],
            schedule=Schedule(weeks=2),
        ),
        "sources": [],
    }
    defaults.update(overrides)
    return ProgramDraft(**defaults)


def _docx_text(data: bytes) -> str:
    doc = Document(io.BytesIO(data))
    paragraphs = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                paragraphs.append(cell.text)
    return "\n".join(paragraphs)


# --------------------------------------------------------------------------- basics
def test_render_returns_bytes_zip():
    data = render_program_draft(_minimal_draft())
    assert isinstance(data, bytes)
    assert data.startswith(b"PK\x03\x04")  # DOCX is a zip archive


def test_render_contains_title_and_kind():
    text = _docx_text(render_program_draft(_minimal_draft()))
    assert "Тестовая программа" in text
    assert "ОБЩЕРАЗВИВАЮЩАЯ" in text


def test_render_professional_development_kind():
    draft = _minimal_draft(kind=ProgramKind.PROFESSIONAL_DEVELOPMENT)
    text = _docx_text(render_program_draft(draft))
    assert "ПОВЫШЕНИЯ КВАЛИФИКАЦИИ" in text


def test_render_names_an_actualization_or_an_adaptation() -> None:
    actualized = _docx_text(render_program_draft(_minimal_draft(mode=ProgramMode.ACTUALIZATION)))
    adapted = _docx_text(render_program_draft(_minimal_draft(mode=ProgramMode.ADAPTATION)))

    assert "актуализация действующей программы" in actualized
    assert "адаптация действующей программы" in adapted


def test_a_five_hundred_hour_programme_renders_without_losing_modules() -> None:
    """The large end of the template range stays a document, not a timeout.

    The model work dominates a real run, but the local render must not grow
    quadratically with the number of modules. The bound is deliberately loose:
    it catches an accidental O(n^2) loop, not a slow CI machine.
    """
    modules = [
        Module(
            id=f"m-{index}",
            title=f"Модуль {index}",
            topics=[
                Topic(
                    id=f"t-{index}.1",
                    title=f"Тема {index}",
                    hours=_hours(self_study=20),
                    outcome_ids=["ro-k-1"],
                )
            ],
        )
        for index in range(1, 26)
    ]
    draft = _minimal_draft(
        curriculum=Curriculum(modules=modules, schedule=Schedule(days=125, hours_per_day=4)),
    )

    started = time.perf_counter()
    data = render_program_draft(draft)
    elapsed = time.perf_counter() - started

    assert elapsed < 10
    text = _docx_text(data)
    assert all(f"Модуль {index}" in text for index in range(1, 26))


def test_render_contains_summary():
    draft = _minimal_draft(
        summary=ProgramSummary(
            goal="Уникальная цель программы",
            tasks=["Задача А", "Задача Б"],
            target_audience="Уникальная аудитория",
            form="Очная",
            issued_document="Сертификат",
            admission_requirements="Требование к входу",
            annotation="Уникальная аннотация",
            relevance="Уникальная актуальность",
            focus="Техническая",
            level="Базовый",
        )
    )
    text = _docx_text(render_program_draft(draft))
    assert "Уникальная цель программы" in text
    assert "Задача А" in text
    assert "Уникальная аудитория" in text
    assert "Требование к входу" in text
    assert "Уникальная аннотация" in text
    assert "Уникальная актуальность" in text
    assert "Техническая" in text
    assert "Базовый" in text


def test_render_competencies_with_all_outcome_types():
    draft = _minimal_draft(
        competencies=[
            Competence(
                id="comp-1",
                code="A/01.6",
                statement="Тестовая компетенция",
                activity="Разработка",
                experience=[Outcome(id="ro-e-1", statement="Опыт А")],
                skills=[Outcome(id="ro-s-1", statement="Умение А")],
                knowledge=[Outcome(id="ro-k-1", statement="Знание А")],
            )
        ]
    )
    text = _docx_text(render_program_draft(draft))
    assert "Тестовая компетенция" in text
    assert "A/01.6" in text
    assert "Опыт А" in text
    assert "Умение А" in text
    assert "Знание А" in text


# --------------------------------------------------------------------------- curriculum
def test_render_curriculum_table():
    draft = _minimal_draft(
        curriculum=Curriculum(
            modules=[
                Module(
                    id="m-1",
                    title="Модуль Альфа",
                    topics=[
                        Topic(
                            id="t-1.1",
                            title="Тема Альфа-1",
                            hours=_hours(lecture=4, practice=6, self_study=2),
                        ),
                        Topic(
                            id="t-1.2",
                            title="Тема Альфа-2",
                            hours=_hours(lecture=2),
                        ),
                    ],
                ),
                Module(
                    id="m-2",
                    title="Модуль Бета",
                    topics=[
                        Topic(
                            id="t-2.1",
                            title="Тема Бета-1",
                            hours=_hours(practice=10),
                        )
                    ],
                ),
            ],
            schedule=Schedule(weeks=3),
        )
    )
    data = render_program_draft(draft)
    doc = Document(io.BytesIO(data))
    tables = doc.tables
    assert len(tables) >= 1

    curriculum_table = tables[0]
    # header + 2 modules + 3 topics + total = 7 rows
    assert len(curriculum_table.rows) == 7

    text = _docx_text(data)
    assert "Модуль Альфа" in text
    assert "Тема Альфа-1" in text
    assert "Итого" in text


def test_render_hours_format_comma_separator():
    draft = _minimal_draft(
        curriculum=Curriculum(
            modules=[
                Module(
                    id="m-1",
                    title="Модуль",
                    topics=[
                        Topic(
                            id="t-1.1",
                            title="Тема",
                            hours=_hours(lecture=0.5, practice=1.5),
                        )
                    ],
                )
            ],
            schedule=Schedule(weeks=1),
        )
    )
    text = _docx_text(render_program_draft(draft))
    assert "0,5" in text
    assert "1,5" in text


def test_render_schedule_fields():
    draft = _minimal_draft(
        curriculum=Curriculum(
            modules=[
                Module(
                    id="m-1",
                    title="Модуль",
                    topics=[Topic(id="t-1.1", title="Тема", hours=_hours(lecture=4))],
                )
            ],
            schedule=Schedule(days=10, hours_per_day=2.5, note="Особое примечание"),
        )
    )
    text = _docx_text(render_program_draft(draft))
    assert "Количество учебных дней" in text
    assert "Режим занятий" in text
    assert "Особое примечание" in text


# --------------------------------------------------------------------------- sources
def test_render_sources_grouped():
    draft = _minimal_draft(
        sources=[
            Source(id="src-1", kind=SourceKind.NORMATIVE_ACT, title="ФЗ-273"),
            Source(
                id="src-2",
                kind=SourceKind.PROFESSIONAL_STANDARD,
                title="Профстандарт 06.001",
                identifier="06.001",
            ),
            Source(
                id="src-3",
                kind=SourceKind.LITERATURE,
                title="Основная книга",
                primary=True,
            ),
            Source(
                id="src-4",
                kind=SourceKind.LITERATURE,
                title="Доп. книга",
                primary=False,
            ),
            Source(
                id="src-5",
                kind=SourceKind.INTERNET_RESOURCE,
                title="Документация",
                url="https://example.test",
            ),
            Source(id="src-6", kind=SourceKind.REFERENCE, title="Справочник"),
        ]
    )
    text = _docx_text(render_program_draft(draft))
    assert "Нормативные акты" in text
    assert "Профессиональные стандарты" in text
    assert "Основная литература" in text
    assert "Дополнительная литература" in text
    assert "Интернет-ресурсы" in text
    assert "Справочные материалы" in text
    assert "ФЗ-273" in text
    assert "https://example.test" in text


def test_the_bibliography_numbers_restart_in_every_group():
    """The "List Number" style keeps one counter for the whole document."""
    draft = _minimal_draft(
        sources=[
            Source(id="src-1", kind=SourceKind.NORMATIVE_ACT, title="ФЗ-273"),
            Source(id="src-2", kind=SourceKind.NORMATIVE_ACT, title="Приказ 499"),
            Source(id="src-3", kind=SourceKind.LITERATURE, title="Книга", primary=True),
        ]
    )

    text = _docx_text(render_program_draft(draft))

    assert "1. ФЗ-273" in text
    assert "2. Приказ 499" in text
    assert "1. Книга" in text


def test_the_document_shows_the_traceability_of_the_programme():
    """A reader has to see which standard a competence rests on."""
    draft = _minimal_draft(
        competencies=[
            Competence(
                id="comp-01",
                statement="Способен анализировать данные",
                standard_source_id="profstd-06.001",
                knowledge=[Outcome(id="ro-k-01", statement="Знает основы статистики")],
            )
        ],
    )

    text = _docx_text(render_program_draft(draft))

    assert "Профессиональный стандарт: profstd-06.001" in text
    assert "ro-k-01. Знает основы статистики" in text


def test_render_no_sources_skips_section():
    text = _docx_text(render_program_draft(_minimal_draft()))
    assert "Литература и источники" not in text


def test_render_marks_a_source_nobody_confirmed():
    """A document the catalog does not hold is marked in the document too."""
    draft = _minimal_draft(sources=[Source(id="src-1", kind=SourceKind.NORMATIVE_ACT, title="ФЗ-273", confirmed=False)])

    text = _docx_text(render_program_draft(draft))

    assert "ФЗ-273" in text
    assert "не подтверждён каталогом источников" in text


# --------------------------------------------------------------------------- pending
def test_render_has_pending_sections_note():
    text = _docx_text(render_program_draft(_minimal_draft()))
    assert "следующем этапе генерации" in text
    assert "Оценочные материалы" in text
    assert "Методические материалы" in text
    assert "Кадровое обеспечение" in text


def test_a_practice_only_programme_has_no_pending_lecture_materials():
    """An internship with no lectures does not owe a lecture plan."""
    draft = _minimal_draft(
        curriculum=Curriculum(
            modules=[
                Module(
                    id="m-1",
                    title="Практика",
                    topics=[
                        Topic(
                            id="t-1.1",
                            title="Стажировка",
                            hours=_hours(self_study=16),
                            outcome_ids=["ro-k-1"],
                        )
                    ],
                )
            ],
            schedule=Schedule(days=4, hours_per_day=4),
        )
    )

    text = _docx_text(render_program_draft(draft))

    assert "Методические материалы" not in text
    assert "Оценочные материалы" in text


# --------------------------------------------------------------------------- overrides
def test_render_accepts_custom_city_and_year():
    text = _docx_text(render_program_draft(_minimal_draft(), city="Санкт-Петербург", year=2030))
    assert "Санкт-Петербург 2030" in text


def test_render_internship_title():
    """A professional programme in internship form says so on the cover."""
    draft = _minimal_draft(kind=ProgramKind.PROFESSIONAL_DEVELOPMENT, format=ProgramFormat.INTERNSHIP)
    text = _docx_text(render_program_draft(draft))
    assert "ДОПОЛНИТЕЛЬНАЯ ПРОФЕССИОНАЛЬНАЯ ПРОГРАММА" in text
    assert "В ФОРМЕ СТАЖИРОВКИ" in text


def test_render_signatures_are_blank_lines():
    """The signature block is left for the methodist to sign by hand."""
    text = _docx_text(render_program_draft(_minimal_draft()))
    assert "Разработчик:" in text
    assert "Руководитель организации:" in text
    assert "______________" in text
