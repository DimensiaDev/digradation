"""Tests for the exports of a programme: the DOCX document and the slides.

The exports are what the methodist hands on, so the tests check that the
document carries the programme and not a placeholder: the title, the
competences, the curriculum, the assessment and the materials.
"""

from __future__ import annotations

import io
import json
import zipfile
from pathlib import Path

from docx import Document
from pptx import Presentation

from app.agentic.models import Materials
from app.render import render_program_document
from app.slides import render_slides

EXAMPLE_PATH = Path(__file__).resolve().parents[1] / "examples" / "document_state_example.json"


def example() -> dict:
    return json.loads(EXAMPLE_PATH.read_text(encoding="utf-8"))


def a_document():
    """The whole example programme, as a document."""
    from app.agentic.state import ProgramDocument

    raw = example()
    return ProgramDocument(
        draft={
            "kind": raw["meta"]["kind"],
            "title": raw["meta"]["title"],
            "summary": raw["summary"],
            "competencies": raw["competencies"],
            "curriculum": raw["curriculum"],
            "sources": raw["sources"],
        },
        assessment=raw["assessment"],
        materials=raw["materials"],
    )


def document_text(document) -> str:
    """Everything the document says, as one string."""
    doc = Document(io.BytesIO(render_program_document(document)))
    parts = [paragraph.text for paragraph in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            parts.extend(cell.text for cell in row.cells)
    return "\n".join(parts)


def test_the_document_carries_the_programme() -> None:
    raw = example()

    text = document_text(a_document())

    assert raw["meta"]["title"] in text
    assert "Оценочные материалы" in text
    assert "Итоговая аттестация" in text
    assert "Тестовые вопросы" in text
    assert "Задания для семинаров" in text
    assert "Методические материалы" in text
    assert "Слайды" in text


def test_the_document_lists_the_competences_and_the_curriculum() -> None:
    raw = example()

    text = document_text(a_document())

    assert raw["competencies"][0]["statement"] in text
    assert raw["curriculum"]["modules"][0]["title"] in text


def test_the_document_is_a_docx_file() -> None:
    content = render_program_document(a_document())

    assert content[:2] == b"PK"
    with zipfile.ZipFile(io.BytesIO(content)) as archive:
        assert "word/document.xml" in archive.namelist()


def test_the_slides_carry_the_decks() -> None:
    raw = example()
    materials = raw["materials"]

    content = render_slides(Materials.model_validate(raw["materials"]), programme_title="Программа")

    assert content[:2] == b"PK"
    prs = Presentation(io.BytesIO(content))
    titles = [slide.shapes.title.text for slide in prs.slides if slide.shapes.title is not None]
    assert "Программа" in titles
    for deck in materials["slide_decks"]:
        assert deck["title"] in titles
    for deck in materials["slide_decks"]:
        for slide in deck["slides"]:
            assert slide["title"] in titles


def test_the_slides_of_one_module_only() -> None:
    raw = example()
    deck = raw["materials"]["slide_decks"][0]

    content = render_slides(
        Materials.model_validate(raw["materials"]), programme_title="Программа", module_ids={deck["topic_id"]}
    )

    prs = Presentation(io.BytesIO(content))
    titles = [slide.shapes.title.text for slide in prs.slides if slide.shapes.title is not None]
    assert deck["title"] in titles
    other = raw["materials"]["slide_decks"][1]
    assert other["title"] not in titles


def test_a_programme_without_slides_still_renders_a_file() -> None:
    content = render_slides(Materials(lecture_plans=[], slide_decks=[]), programme_title="Программа")

    prs = Presentation(io.BytesIO(content))
    assert len(prs.slides) == 1
